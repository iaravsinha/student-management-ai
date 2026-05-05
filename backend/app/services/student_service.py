from io import BytesIO

from fastapi import UploadFile
from fastapi import HTTPException, status
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.faculty import FacultyProfile
from app.models.student import Student
from app.models.timetable import Timetable
from app.models.user import User, UserRole
from app.schemas.student import StudentCreate, StudentImportRowError, StudentImportSummary, StudentUpdate
from app.services import audit_service, department_service


def _ensure_unique_constraints(
    db: Session,
    *,
    email: str | None,
    enrollment_number: str | None,
    roll_number: str | None,
    department: str | None = None,
    exclude_student_id: int | None = None,
) -> None:
  if email is not None:
    query = db.query(Student).filter(Student.email == email)
    if exclude_student_id is not None:
      query = query.filter(Student.id != exclude_student_id)
    if query.first():
      raise HTTPException(
          status_code=status.HTTP_409_CONFLICT,
          detail="Email already in use",
      )

  if enrollment_number is not None:
    query = db.query(Student).filter(Student.enrollment_number == enrollment_number)
    if exclude_student_id is not None:
      query = query.filter(Student.id != exclude_student_id)
    if query.first():
      raise HTTPException(
          status_code=status.HTTP_409_CONFLICT,
          detail="Enrollment number already in use",
      )

  if roll_number is not None:
    query = db.query(Student).filter(Student.roll_number == roll_number)
    if department is not None:
      query = query.filter(Student.department == department)
    if exclude_student_id is not None:
      query = query.filter(Student.id != exclude_student_id)
    if query.first():
      raise HTTPException(
          status_code=status.HTTP_409_CONFLICT,
          detail="Roll number already in use for this department",
      )


def _extract_sequence(value: str, prefix: str) -> int | None:
  if not value.startswith(prefix):
    return None
  suffix = value[len(prefix):]
  return int(suffix) if suffix.isdigit() else None


def _next_student_sequence(db: Session, department_name: str, department_code: str) -> int:
  prefix = f"{department_code}-"
  existing_numbers = (
      db.query(Student.enrollment_number)
      .filter(
          Student.department == department_name,
          Student.enrollment_number.like(f"{prefix}%"),
      )
      .all()
  )
  max_sequence = 0
  for (value,) in existing_numbers:
    sequence = _extract_sequence(value, prefix)
    if sequence is not None:
      max_sequence = max(max_sequence, sequence)
  return max_sequence + 1


def generate_student_identifiers(db: Session, department_name: str, department_code: str) -> tuple[str, str]:
  next_sequence = _next_student_sequence(db, department_name, department_code)
  roll_number = f"{next_sequence:04d}"
  enrollment_number = f"{department_code}-{roll_number}"
  return enrollment_number, roll_number


def get_student_by_user_email(db: Session, email: str) -> Student:
  student = db.query(Student).filter(Student.email == email).first()
  if not student:
    raise HTTPException(
        status_code=status.HTTP_404_NOT_FOUND,
        detail="Student profile not found for this account",
    )
  return student


def ensure_student_access(db: Session, current_user: User, student_id: int) -> Student:
  student = get_student_by_id(db, student_id)
  if current_user.role == UserRole.student and student.email != current_user.email:
    raise HTTPException(
        status_code=status.HTTP_403_FORBIDDEN,
        detail="Students can only access their own record",
    )
  return student


def create_student(db: Session, payload: StudentCreate) -> Student:
  department = department_service.require_department(db, payload.department)
  enrollment_number, roll_number = generate_student_identifiers(db, department.name, department.code)
  _ensure_unique_constraints(
      db,
      email=payload.email,
      enrollment_number=enrollment_number,
      roll_number=roll_number,
      department=department.name,
  )

  student = Student(
      enrollment_number=enrollment_number,
      name=payload.name,
      roll_number=roll_number,
      department=department.name,
      batch_year=payload.batch_year,
      semester=payload.semester,
      email=payload.email,
  )
  db.add(student)
  db.commit()
  db.refresh(student)
  return student


def list_students(
    db: Session,
    *,
    current_user: User,
    page: int,
    page_size: int,
    enrollment_number: str | None = None,
    department: str | None = None,
    batch_year: int | None = None,
    roll_number: str | None = None,
) -> tuple[list[Student], int]:
  query = db.query(Student)
  if current_user.role == UserRole.student:
    query = query.filter(Student.email == current_user.email)
  if current_user.role == UserRole.teacher:
    faculty_profile = db.query(FacultyProfile).filter(FacultyProfile.user_id == current_user.id).first()
    if not faculty_profile:
      raise HTTPException(
          status_code=status.HTTP_404_NOT_FOUND,
          detail="Faculty profile not found for this account",
      )
    query = query.filter(Student.department == faculty_profile.department)
  if enrollment_number:
    query = query.filter(Student.enrollment_number.ilike(f"%{enrollment_number}%"))
  if department:
    query = query.filter(Student.department.ilike(f"%{department}%"))
  if batch_year:
    query = query.filter(Student.batch_year == batch_year)
  if roll_number:
    query = query.filter(Student.roll_number.ilike(f"%{roll_number}%"))

  total = db.query(func.count(Student.id)).select_from(query.subquery()).scalar() or 0
  items = (
      query.order_by(Student.id.desc())
      .offset((page - 1) * page_size)
      .limit(page_size)
      .all()
  )
  return items, total


def get_student_by_id(db: Session, student_id: int) -> Student:
  student = db.query(Student).filter(Student.id == student_id).first()
  if not student:
    raise HTTPException(
        status_code=status.HTTP_404_NOT_FOUND,
        detail="Student not found",
    )
  return student


def update_student(db: Session, student_id: int, payload: StudentUpdate) -> Student:
  student = get_student_by_id(db, student_id)
  updates = payload.model_dump(exclude_unset=True)
  if not updates:
    return student

  updates.pop("enrollment_number", None)
  updates.pop("roll_number", None)
  department = updates.get("department", student.department)
  department_record = department_service.require_department(db, department)
  updates["department"] = department_record.name
  if department_record.name != student.department:
    next_enrollment_number, next_roll_number = generate_student_identifiers(
        db,
        department_record.name,
        department_record.code,
    )
    updates["enrollment_number"] = next_enrollment_number
    updates["roll_number"] = next_roll_number

  _ensure_unique_constraints(
      db,
      email=updates.get("email"),
      enrollment_number=updates.get("enrollment_number"),
      roll_number=updates.get("roll_number"),
      department=updates.get("department", student.department),
      exclude_student_id=student.id,
  )

  for key, value in updates.items():
    setattr(student, key, value)

  db.add(student)
  db.commit()
  db.refresh(student)
  return student


def delete_student(db: Session, student_id: int) -> None:
  student = get_student_by_id(db, student_id)
  db.delete(student)
  db.commit()


def list_students_for_timetable(db: Session, *, current_user: User, timetable_id: int) -> list[Student]:
  timetable = db.query(Timetable).filter(Timetable.id == timetable_id).first()
  if not timetable:
    raise HTTPException(
        status_code=status.HTTP_404_NOT_FOUND,
        detail="Timetable entry not found",
    )

  if current_user.role == UserRole.teacher and timetable.faculty_user_id != current_user.id:
    raise HTTPException(
        status_code=status.HTTP_403_FORBIDDEN,
        detail="Teachers can only access students from their own classes",
    )

  return (
      db.query(Student)
      .filter(
          Student.department == timetable.department,
          Student.batch_year == timetable.batch_year,
          Student.semester == timetable.semester,
      )
      .order_by(Student.name.asc(), Student.id.asc())
      .all()
  )


def _normalize_header(value) -> str:
  return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


async def import_students_from_excel(
    db: Session,
    file: UploadFile,
    *,
    current_user: User,
) -> StudentImportSummary:
  filename = file.filename or ""
  if not filename.lower().endswith(".xlsx"):
    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Upload must be an .xlsx file",
    )

  contents = await file.read()
  if len(contents) > settings.UPLOAD_MAX_BYTES:
    raise HTTPException(
        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
        detail="Uploaded file is too large",
    )

  try:
    workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
  except Exception as exc:  # noqa: BLE001
    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Unable to parse Excel file",
    ) from exc

  sheet = workbook.active
  rows = sheet.iter_rows(values_only=True)
  try:
    headers = [_normalize_header(value) for value in next(rows)]
  except StopIteration as exc:
    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Excel file is empty",
    ) from exc

  required_columns = {"name", "email", "department", "batch_year", "semester"}
  missing_columns = sorted(required_columns - set(headers))
  if missing_columns:
    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail=f"Missing required columns: {', '.join(missing_columns)}",
    )

  inserted = 0
  total_rows = 0
  seen_emails: set[str] = set()
  errors: list[StudentImportRowError] = []

  for row_number, values in enumerate(rows, start=2):
    if not values or all(value in (None, "") for value in values):
      continue
    total_rows += 1
    raw = dict(zip(headers, values, strict=False))
    email = str(raw.get("email") or "").strip().lower()
    if email in seen_emails:
      errors.append(StudentImportRowError(row=row_number, email=email, reason="Duplicate email in uploaded file"))
      continue
    seen_emails.add(email)

    try:
      payload = StudentCreate(
          name=str(raw.get("name") or "").strip(),
          department=str(raw.get("department") or "").strip(),
          batch_year=int(raw.get("batch_year")),
          semester=int(raw.get("semester")),
          email=email,
      )
      create_student(db, payload)
      inserted += 1
    except (TypeError, ValueError, ValidationError) as exc:
      errors.append(StudentImportRowError(row=row_number, email=email or None, reason=str(exc)))
    except HTTPException as exc:
      errors.append(StudentImportRowError(row=row_number, email=email or None, reason=str(exc.detail)))

  summary = StudentImportSummary(
      inserted=inserted,
      skipped=len(errors),
      total_rows=total_rows,
      errors=errors,
  )
  audit_service.log_action(
      db,
      actor=current_user,
      action="student:import",
      entity_type="student",
      status="success" if not errors else "partial_success",
      meta=summary.model_dump(),
  )
  return summary

